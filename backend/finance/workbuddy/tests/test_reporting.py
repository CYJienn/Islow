"""Offline reporting regressions: frozen arithmetic, provenance and safe snapshots."""
from copy import deepcopy
from decimal import Decimal, localcontext
from html import escape
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from openpyxl import load_workbook
from shared import contracts
from reporting import DISCLAIMER, SHEETS, generate, summarize


def result(identity='B1', invoice_id='I1', amount='100.10', currency='MYR', status='FAIL'):
    invoice = {'invoice_id': invoice_id, 'invoice_number': 'INV-' + invoice_id,
               'vendor_id': 'V1', 'vendor_name': 'Example Vendor', 'amount': amount,
               'currency': currency, 'invoice_date': '2026-01-10', 'source_ref': 'invoice:' + invoice_id}
    findings = {
        'intake': contracts.finding('INTAKE.AMOUNT', 'PASS', [invoice['source_ref']], amount,
                                   'Finite explicit amount', 'Header checked.', 'Retain evidence.'),
        'matching': contracts.finding('MATCH.PRICE', status, ['line:1', 'po:1', 'receipt:1'],
                                     {'unit_price': '10.01', 'po_id': 'P1', 'po_line_id': 'L1'},
                                     {'unit_price': '10.00', 'currency': currency}, 'Exact price comparison.',
                                     'Review price discrepancy.', evidence='MISSING' if status == 'UNKNOWN' else 'CONFIRMED'),
        'fraud': contracts.finding('FRAUD.DUPLICATE', 'PASS', [invoice['source_ref']], 0, 0,
                                  'No duplicate in supplied records.', 'Retain evidence.'),
        'accounting': contracts.finding('ACCOUNTING.GL', 'PASS', ['gl:1'], {'gl_account': '6100'},
                                       'Active GL', 'Explicit GL checked.', 'Retain evidence.'),
    }
    return {'business_id': identity, 'invoice': invoice,
            'domains': {domain: {'domain': domain, 'business_id': identity, 'findings': [finding], 'errors': []}
                        for domain, finding in findings.items()},
            'status': status, 'flagged': status == 'FAIL', 'unresolved': status == 'UNKNOWN'}


def report(results=None, kind='SUPPLIED'):
    values = [result()] if results is None else results
    return {'report_id': 'IA-TEST', 'version': 1, 'created_at': '2026-10-10T00:00:00Z',
            'dataset_id': 'DS1', 'batch_id': 'batch1', 'dataset_kind': kind, 'results': values,
            'metrics': summarize(values, kind), 'source_manifest': [{'source_ref': 'export:1', 'filename': 'source.xlsx'}],
            'processing_errors': [], 'processing_seconds': 1.25, 'review_seconds': None,
            'policy': {}, 'demo_expected': {},
            'reference_tables': {
                'invoice_lines': [{'invoice_id': 'I1', 'po_id': 'P1', 'po_line_id': 'L1',
                                   'quantity': '10', 'unit_price': '10.01', 'currency': 'MYR', 'source_ref': 'line:1'}],
                'po_lines': [{'po_id': 'P1', 'po_line_id': 'L1', 'unit_price': '10.00',
                              'currency': 'MYR', 'source_ref': 'po:1'}],
                'receipts': [{'receipt_id': 'R1', 'po_id': 'P1', 'po_line_id': 'L1',
                              'quantity': '10', 'source_ref': 'receipt:1'}],
                'vendors': [{'vendor_id': 'V1', 'vendor_name': 'Sourced Vendor', 'source_ref': 'vendor:1'}],
                'gl_accounts': [{'gl_account': '6100', 'active': True, 'source_ref': 'gl:1'}],
            }}


class SummarizeTests(unittest.TestCase):
    def test_unique_business_counts_not_domain_findings(self):
        first = result()
        first['domains']['fraud']['findings'].append(contracts.finding(
            'FRAUD.SIMILAR', 'FAIL', ['invoice:I1'], 100, 'Review', 'Candidate.', 'Review.',
            evidence='INFERRED_CANDIDATE'))
        metrics = summarize([first, deepcopy(first), result('B2', 'I2', '50', status='PASS')], 'SUPPLIED')
        self.assertEqual({k: metrics[k] for k in ('total', 'covered', 'flagged', 'unresolved')},
                         {'total': 2, 'covered': 2, 'flagged': 1, 'unresolved': 0})
        self.assertEqual(metrics['flagged_exposure_by_currency'], {'MYR': '100.10'})
        self.assertEqual(metrics['flagged_exposure_by_currency'], metrics['supplied_flagged_exposure_by_currency'])
        contracts.canonical(metrics)

    def test_unique_once_per_currency_no_fx_netting(self):
        metrics = summarize([result(amount='0.1'), result('B2', 'I2', '0.2'),
                             result('B1', 'I1', '25', currency='USD')], 'SUPPLIED')
        self.assertEqual(metrics['total'], 2)
        self.assertEqual(metrics['flagged_exposure_by_currency'], {'MYR': '0.3', 'USD': '25'})

    def test_decimal_arithmetic_not_float_or_global_context(self):
        large = '123456789012345678901234567890.000000000000000000000000000001'
        with localcontext() as context:
            context.prec = 6
            metrics = summarize([result(amount=large), result('B2', 'I2', '0.000000000000000000000000000009')], 'SUPPLIED')
            self.assertEqual(context.prec, 6)
        self.assertEqual(metrics['flagged_exposure_by_currency']['MYR'],
                         '123456789012345678901234567890.000000000000000000000000000010')

    def test_all_demo_supplied_headline_zero(self):
        metrics = summarize([result()], 'SYNTHETIC_DEMO')
        self.assertEqual(metrics['flagged_exposure_by_currency'], {'MYR': '0'})
        self.assertEqual(metrics['supplied_flagged_exposure_by_currency'], {'MYR': '0'})
        self.assertEqual(metrics['demo_flagged_exposure_by_currency'], {'MYR': '100.10'})

    def test_mixed_demo_markers_and_duplicate_identity_veto(self):
        demo = result('B2', 'I2', '900')
        demo['invoice']['dataset_kind'] = 'SYNTHETIC_DEMO'
        other = result('B3', 'I3', '80', currency='USD')
        other['dataset_kind'] = 'SYNTHETIC_DEMO'
        metrics = summarize([result(), demo, other], 'SUPPLIED')
        self.assertEqual(metrics['flagged_exposure_by_currency'], {'MYR': '100.10', 'USD': '0'})
        self.assertEqual(metrics['demo_flagged_exposure_by_currency'], {'MYR': '900', 'USD': '80'})
        duplicated = result()
        duplicated['dataset_kind'] = 'SYNTHETIC_DEMO'
        metrics = summarize([result(), duplicated], 'SUPPLIED')
        self.assertEqual(metrics['flagged_exposure_by_currency'], {'MYR': '0'})
        self.assertEqual(metrics['demo_flagged_exposure_by_currency'], {'MYR': '100.10'})

    def test_unknown_flagged_amount_and_currency_are_not_zero_facts(self):
        for value in (None, '', True, 'NaN', 'Infinity', 'invalid'):
            with self.subTest(amount=value):
                metrics = summarize([result(amount=value)], 'SUPPLIED')
                self.assertEqual(metrics['unknown_amount_flagged'], 1)
                self.assertEqual(metrics['flagged_exposure_by_currency'], {})
        for currency in (None, '', 'myr', '=USD'):
            metrics = summarize([result(currency=currency)], 'SUPPLIED')
            self.assertEqual(metrics['unknown_amount_flagged'], 1)
            self.assertEqual(metrics['flagged_exposure_by_currency'], {})
        self.assertEqual(summarize([result(amount=None, status='UNKNOWN')], 'SUPPLIED')['unknown_amount_flagged'], 0)

    def test_flags_and_unknown_derived_from_findings_fail_closed(self):
        first = result(status='UNKNOWN')
        first.update(status='PASS', flagged=False, unresolved=False)
        self.assertEqual(summarize([first], 'SUPPLIED')['unresolved'], 1)
        first['domains']['intake']['findings'][0]['status'] = 'FAIL'
        self.assertEqual(summarize([first], 'SUPPLIED')['flagged'], 1)
        self.assertEqual(summarize([first], 'SUPPLIED')['covered'], 1)
        del first['domains']['accounting']
        self.assertEqual(summarize([first], 'SUPPLIED')['covered'], 0)
        self.assertEqual(summarize([first], 'SUPPLIED')['unresolved'], 1)

    def test_processing_errors_and_empty_findings_are_not_coverage(self):
        first = result(status='PASS')
        for mutation in ('error', 'empty'):
            value = deepcopy(first)
            if mutation == 'error':
                value['domains']['matching']['errors'] = ['SOURCE_UNAVAILABLE']
            else:
                value['domains']['matching']['findings'] = []
            metrics = summarize([value], 'SUPPLIED')
            self.assertEqual(metrics['covered'], 0)
            self.assertEqual(metrics['unresolved'], 1)

    def test_conflicting_duplicate_amount_refuses_guessing(self):
        with self.assertRaisesRegex(ValueError, 'CONFLICTING_BUSINESS_INVOICE_AMOUNT'):
            summarize([result(), result(amount='200')], 'SUPPLIED')

    def test_empty_and_invalid_identity_or_kind(self):
        metrics = summarize([], 'SUPPLIED')
        self.assertTrue(all(metrics[key] == 0 for key in ('total', 'covered', 'flagged', 'unresolved', 'unknown_amount_flagged')))
        with self.assertRaisesRegex(ValueError, 'INVALID_DATASET_KIND'):
            summarize([], 'UNKNOWN')
        value = result()
        value['business_id'] = ''
        with self.assertRaisesRegex(ValueError, 'MISSING_STABLE_BUSINESS_ID'):
            summarize([value], 'SUPPLIED')


class GenerateTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.base = Path(self.directory.name)

    def generate(self, content=None, name='snapshot'):
        content = report() if content is None else content
        original = deepcopy(content)
        # Any network access or automatic installation is an immediate regression.
        with patch('socket.socket', side_effect=AssertionError('NETWORK_FORBIDDEN')), \
                patch('subprocess.check_call', side_effect=AssertionError('INSTALL_FORBIDDEN')):
            paths = generate(content, self.base / name)
        self.assertEqual(content, original)
        return paths, json.loads(Path(paths['data']).read_text(encoding='utf-8'))

    def workbook(self, paths, data_only=False):
        workbook = load_workbook(paths['workbook'], data_only=data_only)
        self.addCleanup(workbook.close)
        return workbook

    def summary(self, paths):
        sheet = self.workbook(paths, data_only=True)['Summary']
        return {row[0]: row[1:] for row in sheet.iter_rows(min_row=4, values_only=True)}

    def artifact_text(self, paths):
        text = Path(paths['dashboard']).read_text(encoding='utf-8') + Path(paths['data']).read_text(encoding='utf-8')
        with ZipFile(paths['workbook']) as archive:
            text += ''.join(archive.read(name).decode('utf-8') for name in archive.namelist() if name.endswith('.xml'))
        return text

    def test_absolute_runner_api_and_five_frozen_sheets(self):
        paths, data = self.generate()
        self.assertEqual(set(paths), {'dashboard', 'workbook', 'data'})
        self.assertTrue(all(Path(path).is_absolute() and Path(path).is_file() for path in paths.values()))
        self.assertEqual({p.name for p in (self.base / 'snapshot').iterdir()}, {'dashboard.html', 'assurance.xlsx', 'report.json'})
        book = self.workbook(paths)
        self.assertEqual(tuple(book.sheetnames), SHEETS)
        for sheet in book:
            self.assertEqual(sheet.freeze_panes, 'B4')
            self.assertTrue(sheet.auto_filter.ref.startswith('A3:'))
            self.assertFalse(sheet.merged_cells.ranges)
            self.assertTrue(all(cell.data_type != 'f' for row in sheet for cell in row))
            self.assertTrue(all(cell.alignment.wrap_text for row in sheet.iter_rows(min_row=3) for cell in row))
            self.assertGreater(sheet.column_dimensions['A'].width, 20)
        self.assertEqual(data['metrics']['total'], 1)
        summary = self.summary(paths)
        self.assertEqual(summary['Total'][0], 1)
        self.assertEqual(summary['Supplied flagged exposure (headline)'][0], 100.1)
        self.assertEqual(summary['Processing Seconds'][0], 1.25)
        self.assertEqual(summary['Review Seconds'][0], 'NOT_MEASURED')
        self.assertEqual(data['review_seconds'], None)
        self.assertEqual(data['review_time_status'], 'NOT_MEASURED')
        self.assertEqual(data['cost_status'], 'NOT_CALCULATED')
        self.assertNotIn('cost_estimate', data)

    def test_recompute_stale_metrics_without_mutating_input(self):
        content = report()
        content['metrics'].update(total=999, flagged_exposure_by_currency={'MYR': '999999'})
        paths, data = self.generate(content)
        self.assertEqual(data['metrics']['total'], 1)
        self.assertEqual(data['metrics']['flagged_exposure_by_currency'], {'MYR': '100.10'})
        self.assertNotIn('999999', self.artifact_text(paths))

    def test_html_native_drilldown_all_findings_and_exact_source_records(self):
        paths, data = self.generate()
        html = Path(paths['dashboard']).read_text(encoding='utf-8')
        self.assertIn('color-scheme:light', html)
        self.assertNotIn('<script', html)
        self.assertNotIn('src="http', html)
        self.assertIn('<details class="invoice">', html)
        self.assertIn('Sourced Vendor', html)
        self.assertIn('100.10', html)
        self.assertIn('MYR', html)
        for domain in contracts.DOMAINS:
            self.assertIn('>' + domain.title() + '</summary>', html)
        for _, envelope in data['results'][0]['domains'].items():
            for finding in envelope['findings']:
                for key in ('rule_id', 'status', 'explanation', 'recommended_action', 'evidence_class'):
                    self.assertIn(escape(finding[key]), html)
        for value in ('po:1', 'receipt:1', 'po_lines', 'receipts', 'P1', 'L1', 'R1'):
            self.assertIn(value, html)
        self.assertIn('Observed', html)
        self.assertIn('Expected', html)
        self.assertIn(DISCLAIMER, html)
        evidence = list(self.workbook(paths)['Evidence References'].iter_rows(min_row=4, values_only=True))
        self.assertTrue(any(row[4] == 'receipt:1' and row[5] == 'receipts' and 'R1' in row[7] for row in evidence))
        self.assertEqual(len(list(self.workbook(paths)['Exception Details'].iter_rows(min_row=4))), 4)

    def test_without_reference_tables_preserves_refs_not_inferred_links(self):
        content = report()
        del content['reference_tables']
        paths, data = self.generate(content)
        html = Path(paths['dashboard']).read_text(encoding='utf-8')
        self.assertIn('SOURCE_ROW_NOT_AVAILABLE', html)
        self.assertIn('receipt:1', html)
        self.assertNotIn('"receipt_id": "R1"', html)
        self.assertNotIn('reference_tables', data)

    def test_recursive_redaction_public_finding_and_no_raw_payload(self):
        content = report()
        invoice = content['results'][0]['invoice']
        invoice.update(bank_account='PRIVATE-INVOICE-111', account_number='PRIVATE-ACCOUNT-222',
                       nested={'credentials': {'password': 'PRIVATE-PASS-333'}, 'api_key': 'PRIVATE-API-444'},
                       raw={'note': 'RAW-INSTRUCTIONS-NOT-FOR-REPORT', 'bank_account': 'PRIVATE-RAW-555'})
        content['reference_tables']['vendors'][0].update(iban='PRIVATE-IBAN-666', swift='PRIVATE-SWIFT-777')
        content['source_manifest'][0]['access_token'] = 'PRIVATE-TOKEN-888'
        content['processing_errors'] = [{'error': 'failed PRIVATE-PASS-333 PRIVATE-INVOICE-111'}]
        bank = contracts.finding('FRAUD.BANK_MATCH', 'FAIL', ['vendor:1'], 'PRIVATE-BANK-999',
                                 'PRIVATE-EXPECTED-000', 'Compare PRIVATE-BANK-999 against PRIVATE-EXPECTED-000.',
                                 'Never print PRIVATE-INVOICE-111.')
        content['results'][0]['domains']['fraud']['findings'].append(bank)
        paths, data = self.generate(content)
        combined = self.artifact_text(paths)
        for secret in ('PRIVATE-INVOICE-111', 'PRIVATE-ACCOUNT-222', 'PRIVATE-PASS-333', 'PRIVATE-API-444',
                       'PRIVATE-RAW-555', 'PRIVATE-IBAN-666', 'PRIVATE-SWIFT-777', 'PRIVATE-TOKEN-888',
                       'PRIVATE-BANK-999', 'PRIVATE-EXPECTED-000', 'RAW-INSTRUCTIONS-NOT-FOR-REPORT'):
            self.assertNotIn(secret, combined)
        self.assertNotIn('raw', data['results'][0]['invoice'])
        self.assertEqual(data['results'][0]['invoice']['bank_account'], '[REDACTED]')
        public = data['results'][0]['domains']['fraud']['findings'][-1]
        self.assertEqual(public['observed_value'], '[BANK COMPARISON REDACTED]')
        self.assertEqual(public['expected_value'], 'Independently verified vendor bank details; values withheld')
        self.assertEqual(data['reference_tables']['gl_accounts'][0]['gl_account'], '6100')

    def test_html_escape_formula_injection_and_invalid_xml(self):
        content = report()
        payload = '<script>alert("unsafe")</script>&\x00\x01\ud800\ufffe'
        invoice = content['results'][0]['invoice']
        invoice['vendor_name'] = payload
        invoice['notes'] = '=HYPERLINK("https://invalid.example","unsafe")'
        content['reference_tables']['vendors'][0]['vendor_name'] = payload
        for table, prefix in zip(('invoice_lines', 'po_lines', 'receipts', 'vendors'), ('=', '+', '-', '@')):
            content['reference_tables'][table][0]['note'] = prefix + 'SOURCE-PAYLOAD'
        for domain in content['results'][0]['domains'].values():
            domain['findings'][0].update(observed_value='=1+1', expected_value='+1+1',
                                         explanation='-1+1', recommended_action='@SUM(1,1)')
        content['report_id'] = '=REPORT'
        paths, data = self.generate(content)
        html = Path(paths['dashboard']).read_text(encoding='utf-8')
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('\x00', self.artifact_text(paths))
        self.assertEqual(data['results'][0]['invoice']['vendor_name'], '<script>alert("unsafe")</script>&')
        book = self.workbook(paths)
        literals = [cell for sheet in book for row in sheet for cell in row
                    if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@'))]
        self.assertGreaterEqual(len(literals), 17)
        self.assertTrue(all(cell.data_type == 's' for cell in literals))
        with ZipFile(paths['workbook']) as archive:
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    self.assertNotIn(b'<f>', archive.read(name))

    def test_large_exact_amount_retained_not_excel_rounded(self):
        amount = '123456789012345678901234567890.000000000000000001'
        paths, data = self.generate(report([result(amount=amount)]))
        self.assertEqual(data['metrics']['flagged_exposure_by_currency']['MYR'], amount)
        self.assertEqual(self.summary(paths)['Supplied flagged exposure (headline)'][0], amount)
        row = next(self.workbook(paths, True)['Invoice Results'].iter_rows(min_row=4, values_only=True))
        self.assertEqual(row[10], amount)
        self.assertEqual(row[12], amount)

    def test_demo_expected_rule_mapping_and_supplied_zero_headline(self):
        content = report(kind='SYNTHETIC_DEMO')
        content['demo_expected'] = {'I1': {'MATCH.PRICE': 'FAIL', 'FRAUD.DUPLICATE': 'FAIL'},
                                    'MISSING': {'MATCH.QUANTITY': 'PASS'}}
        paths, data = self.generate(content)
        self.assertEqual(data['metrics']['flagged_exposure_by_currency'], {'MYR': '0'})
        rows = list(self.workbook(paths)['Demo Cases'].iter_rows(min_row=4, values_only=True))
        self.assertEqual(rows[0][:5], ('I1', 'MATCH.PRICE', 'FAIL', 'FAIL', 'PASS'))
        self.assertEqual(rows[1][4], 'FAIL')
        self.assertEqual(rows[2][3:5], ('UNKNOWN', 'UNKNOWN'))
        summary = self.summary(paths)
        self.assertEqual(summary['Supplied flagged exposure (headline)'][0], 0)
        self.assertEqual(summary['Synthetic demo flagged exposure (separate)'][0], 100.1)
        self.assertIn('excluded from supplied headline exposure', Path(paths['dashboard']).read_text(encoding='utf-8'))

    def test_missing_measurement_no_fabricated_cost_or_time(self):
        content = report()
        content['policy'] = {'baseline_minutes_per_invoice': '10', 'hourly_cost': '60', 'cost_currency': 'MYR'}
        content['cost_estimate'] = {'estimated_cost_savings': '999'}
        paths, data = self.generate(content)
        self.assertIsNone(data['review_seconds'])
        self.assertNotIn('cost_estimate', data)
        self.assertEqual(data['cost_status'], 'NOT_CALCULATED')
        for value in (None, -1, True, 'NaN', 'Infinity', 'unknown'):
            content['review_seconds'] = value
            _, data = self.generate(content, 'invalid-review-' + str(value))
            self.assertIsNone(data['review_seconds'])
            self.assertEqual(data['review_time_status'], 'NOT_MEASURED')
            self.assertNotIn('cost_estimate', data)

    def test_cost_requires_supplied_baseline_hourly_currency_and_measurements(self):
        content = report()
        content.update(processing_seconds=60, review_seconds=120)
        content['policy'] = {'baseline_minutes_per_invoice': '10', 'hourly_cost': '60', 'cost_currency': 'MYR'}
        paths, data = self.generate(content)
        estimate = data['cost_estimate']
        self.assertEqual(estimate['baseline_seconds'], '600')
        self.assertEqual(estimate['measured_total_seconds'], '180')
        self.assertEqual(estimate['baseline_cost'], '10.00')
        self.assertEqual(estimate['measured_time_cost'], '3.00')
        self.assertEqual(estimate['measured_time_cost_per_invoice'], '3.0000')
        self.assertEqual(estimate['estimated_cost_savings'], '7.00')
        self.assertEqual(self.summary(paths)['Measured Time Cost Per Invoice'][0], 3)
        for key in ('baseline_minutes_per_invoice', 'hourly_cost', 'cost_currency'):
            bad = deepcopy(content)
            del bad['policy'][key]
            _, data = self.generate(bad, 'missing-' + key)
            self.assertNotIn('cost_estimate', data)
        bad = deepcopy(content)
        bad['processing_seconds'] = None
        _, data = self.generate(bad, 'unmeasured-processing')
        self.assertNotIn('cost_estimate', data)

    def test_zero_invoices_safe_division_and_negative_savings_preserved(self):
        content = report([])
        content.update(review_seconds=0, processing_seconds=0)
        content['policy'] = {'baseline_minutes_per_invoice': '0', 'hourly_cost': '60', 'cost_currency': 'MYR'}
        _, data = self.generate(content)
        self.assertNotIn('cost_estimate', data)
        content['results'] = [result()]
        content.update(processing_seconds=0, review_seconds=60)
        _, data = self.generate(content, 'negative-savings')
        self.assertEqual(data['cost_estimate']['estimated_cost_savings'], '-1.00')
        self.assertEqual(data['cost_estimate']['measured_time_cost_per_invoice'], '1.0000')

    def test_refuses_existing_directory_and_keeps_snapshot_bytes(self):
        paths, _ = self.generate()
        before = {key: Path(path).read_bytes() for key, path in paths.items()}
        with self.assertRaisesRegex(FileExistsError, 'SNAPSHOT_DIRECTORY_ALREADY_EXISTS'):
            generate(report(), self.base / 'snapshot')
        self.assertEqual(before, {key: Path(path).read_bytes() for key, path in paths.items()})
        empty = self.base / 'empty'
        empty.mkdir()
        with self.assertRaises(FileExistsError):
            generate(report(), empty)
        self.assertEqual(list(empty.iterdir()), [])

    def test_invalid_report_does_not_create_snapshot_directory(self):
        content = report()
        content['results'][0]['business_id'] = ''
        with self.assertRaises(ValueError):
            generate(content, self.base / 'invalid')
        self.assertFalse((self.base / 'invalid').exists())

    def test_runner_real_envelopes_header_only_matching_unknown(self):
        from checks import evaluate
        from runner import assurance
        from test_checks import make_batch
        batch = make_batch()
        batch['tables']['invoice_lines'] = []
        invoice = batch['tables']['invoices'][0]
        assembled = assurance(batch, invoice, {domain: evaluate(batch, invoice, domain) for domain in contracts.DOMAINS})
        content = report([assembled])
        content['reference_tables'] = batch['tables']
        paths, data = self.generate(content)
        self.assertEqual(data['metrics']['unresolved'], 1)
        self.assertEqual(data['results'][0]['status'], 'UNKNOWN')
        findings = data['results'][0]['domains']['matching']['findings']
        self.assertTrue(all(finding['status'] == 'UNKNOWN' for finding in findings))
        self.assertIn('header-only matching is not sufficient', Path(paths['dashboard']).read_text(encoding='utf-8'))
        rows = list(self.workbook(paths)['Exception Details'].iter_rows(min_row=4, values_only=True))
        self.assertTrue(all(row[4] == 'UNKNOWN' for row in rows if row[2] == 'matching'))


if __name__ == '__main__':
    unittest.main()
