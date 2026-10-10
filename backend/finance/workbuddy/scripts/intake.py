"""Explicit workbook mapping; no source instructions or macros are executed."""
import argparse
import json
import sys
from pathlib import Path
from datetime import date, datetime
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from shared.contracts import TABLES, now, file_hash, validate, canonical, business_id


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def check_batch(batch):
    validate(batch, 'batch')
    for invoice in batch['tables']['invoices']:
        business_id(batch, invoice)
    return batch


def normalize(source, mapping, dataset_id, policy=None):
    from openpyxl import load_workbook
    source = Path(source)
    if source.suffix.lower() != '.xlsx':
        raise ValueError('Only .xlsx exports are supported; macros and legacy workbook execution are not allowed')
    sha = file_hash(source)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    config = json.loads(Path(mapping).read_text(encoding='utf-8'))
    tables = {t: [] for t in TABLES}
    errors = []
    try:
        for table, spec in config['tables'].items():
            if table not in TABLES:
                raise ValueError('Unknown canonical table: ' + table)
            sheet_name = spec['sheet']
            if sheet_name not in workbook.sheetnames:
                if table == 'invoices':
                    raise ValueError('Required invoice header sheet is missing: ' + sheet_name)
                continue
            columns = spec.get('columns', {})
            if not columns:
                raise ValueError('Explicit column mappings required for sheet: ' + sheet_name)
            ws = workbook[sheet_name]
            rows = iter(ws.iter_rows(values_only=True))
            headers = [str(v).strip() if v is not None else '' for v in next(rows)]
            if len(headers) != len(set(headers)):
                raise ValueError('Duplicate or blank column names in sheet: ' + sheet_name)
            for field, header in columns.items():
                if header not in headers:
                    raise ValueError('Mapped column absent: ' + sheet_name + '/' + header)
                if field in ('source_ref', 'raw'):
                    raise ValueError('Reserved provenance field cannot be mapped')
            for row_number, values in enumerate(rows, start=2):
                if all(v is None for v in values):
                    continue
                raw = {h: scalar(v) for h, v in zip(headers, values)}
                ref = 'sha256:' + sha + '/' + sheet_name + '/row:' + str(row_number)
                record = {k: raw[h] for k, h in columns.items()}
                record.update(source_ref=ref, raw=raw)
                for key, value in list(record.items()):
                    if isinstance(value, str) and value.startswith('='):
                        record[key] = None
                        errors.append('UNEVALUATED_FORMULA:' + ref + '/' + key)
                tables[table].append(record)
    finally:
        workbook.close()
    batch = {'schema_version': '1.0', 'batch_id': 'batch-' + sha[:16], 'dataset_id': dataset_id, 'dataset_kind': 'SUPPLIED', 'retrieved_at': now(), 'source_manifest': [{'document_id': 'sha256:' + sha, 'sha256': sha, 'source_name': source.name, 'source_kind': 'SUPPLIED_SYNTHETIC_EXPORT', 'trust': 'DATA_ONLY', 'path': str(source.resolve())}], 'tables': tables, 'policy': policy or {}, 'retrieval_errors': errors}
    return check_batch(batch)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    p = commands.add_parser('normalize')
    p.add_argument('--source', required=True)
    p.add_argument('--mapping', required=True)
    p.add_argument('--dataset-id', required=True)
    p.add_argument('--policy')
    p.add_argument('--output', required=True)
    p = commands.add_parser('validate')
    p.add_argument('--input', required=True)
    args = parser.parse_args()
    try:
        if args.command == 'normalize':
            policy = json.loads(Path(args.policy).read_text(encoding='utf-8')) if args.policy else {}
            batch = normalize(args.source, args.mapping, args.dataset_id, policy)
            Path(args.output).write_text(json.dumps(batch, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
            print(canonical({'status': 'NORMALIZED', 'batch_id': batch['batch_id'], 'records': len(batch['tables']['invoices']), 'errors': batch['retrieval_errors']}))
        else:
            batch = check_batch(json.loads(Path(args.input).read_text(encoding='utf-8')))
            print(canonical({'status': 'VALIDATED', 'records': len(batch['tables']['invoices'])}))
    except Exception as exc:
        # No raw data or credentials in diagnostic text.
        print(canonical({'status': 'RETRIEVAL_OR_INPUT_FAILURE', 'error_type': type(exc).__name__, 'error': str(exc)[:350]}))
        raise SystemExit(2)

if __name__ == '__main__':
    main()
