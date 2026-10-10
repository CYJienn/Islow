"""Offline, redacted approval snapshots; no live formulas or payment authority."""
from collections import defaultdict
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP, localcontext
from html import escape
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from shared import contracts

SHEETS = ('Summary', 'Invoice Results', 'Exception Details', 'Evidence References', 'Demo Cases')
DISCLAIMER = ('Flagged exposure is suspected risk, not confirmed fraud or prevented loss. '
              'Distribution approval does not authorize payment or clear exceptions.')
XML_INVALID = re.compile(r'[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]')
SENSITIVE = re.compile(
    r'bank|iban|swift|password|passwd|secret|token|credential|api[_-]?key|access[_-]?key|'
    r'authorization|private[_-]?key|routing[_-]?(?:number|no)|(?:^|[_-])pin(?:$|[_-])|'
    r'(?:^|[_-])accounts?(?:$|[_-])|'
    r'account(?:number|no|id)|acct(?:[_-]?(?:number|no))?', re.I)
RAW_KEYS = {'raw', 'raw_data', 'raw_content', 'raw_text', 'raw_records', 'raw_payload',
            'raw_source', 'raw_input', 'raw_inputs', 'raw_document', 'raw_documents'}


def _text(value):
    return XML_INVALID.sub('', str(value))


def _sensitive(key):
    # GL codes and accounting values are business evidence, not bank accounts.
    return key.lower() not in {'gl_account', 'gl_accounts', 'accounting_currency',
                               'accounting_amount'} and bool(SENSITIVE.search(key))


def _sanitize(value):
    secrets = set()

    def collect(item, protected=False):
        if isinstance(item, dict):
            for key, child in item.items():
                collect(child, protected or _sensitive(str(key)))
            if str(item.get('rule_id', '')).startswith('FRAUD.BANK'):
                collect(item.get('observed_value'), True)
                collect(item.get('expected_value'), True)
        elif isinstance(item, (list, tuple)):
            for child in item:
                collect(child, protected)
        elif protected and isinstance(item, (str, int, Decimal)) and not isinstance(item, bool):
            token = _text(item)
            if token and not token.startswith('[REDACTED]'):
                secrets.add(token)

    collect(value)
    tokens = sorted(secrets, key=len, reverse=True)

    def clean_string(item):
        item = _text(item)
        for token in tokens:
            item = item.replace(token, '[REDACTED]')
        return item

    def clean(item):
        if isinstance(item, dict):
            if 'rule_id' in item and 'status' in item:
                item = contracts.public_finding(item)
            return {clean_string(key): '[REDACTED]' if _sensitive(str(key)) else clean(child)
                    for key, child in item.items() if str(key).lower() not in RAW_KEYS}
        if isinstance(item, (list, tuple)):
            return [clean(child) for child in item]
        if isinstance(item, str):
            return clean_string(item)
        if isinstance(item, Decimal):
            return format(item, 'f') if item.is_finite() else None
        if isinstance(item, float) and contracts.number(item) is None:
            return None
        if item is None or isinstance(item, (bool, int, float)):
            return item
        raise TypeError('UNSUPPORTED_REPORT_VALUE')

    return clean(value)


def _findings(result):
    for domain, envelope in result.get('domains', {}).items():
        if isinstance(envelope, dict):
            for finding in envelope.get('findings', []):
                yield domain, finding


def _classification(result):
    domains = result.get('domains', {})
    covered = all(isinstance(domains.get(domain), dict)
                  and domains[domain].get('findings')
                  and not domains[domain].get('errors') for domain in contracts.DOMAINS)
    findings = [finding for _, finding in _findings(result)]
    flagged = result.get('flagged') is True or result.get('status') == 'FAIL' or any(
        f.get('status') == 'FAIL' for f in findings)
    unresolved = not covered or result.get('unresolved') is True or result.get('status') == 'UNKNOWN' or any(
        f.get('status') not in ('PASS', 'FAIL', 'NOT_APPLICABLE') for f in findings)
    return covered, flagged, unresolved


def _is_demo(result, dataset_kind):
    return 'SYNTHETIC_DEMO' in (dataset_kind, result.get('dataset_kind'),
                                result.get('invoice', {}).get('dataset_kind'))


def _precision(numbers, count=1):
    numbers = [n for n in numbers if n is not None and n.is_finite()]
    if not numbers:
        return 50
    return max(50, max(n.adjusted() for n in numbers)
               - min(n.as_tuple().exponent for n in numbers) + len(str(count)) + 20)


def summarize(results, dataset_kind):
    """Count unique business invoices and freeze per-currency Decimal exposure."""
    if dataset_kind not in ('SUPPLIED', 'SYNTHETIC_DEMO'):
        raise ValueError('INVALID_DATASET_KIND')
    businesses = defaultdict(list)
    for result in results:
        identity = result.get('business_id')
        if not isinstance(identity, str) or not identity.strip():
            raise ValueError('MISSING_STABLE_BUSINESS_ID')
        businesses[identity].append(result)
    metrics = {'total': len(businesses), 'covered': 0, 'flagged': 0, 'unresolved': 0,
               'flagged_exposure_by_currency': {}, 'supplied_flagged_exposure_by_currency': {},
               'demo_flagged_exposure_by_currency': {}, 'unknown_amount_flagged': 0}
    supplied, demo = defaultdict(list), defaultdict(list)
    currencies = set()
    for records in businesses.values():
        states = [_classification(r) for r in records]
        metrics['covered'] += int(all(s[0] for s in states))
        flagged = any(s[1] for s in states)
        metrics['flagged'] += int(flagged)
        metrics['unresolved'] += int(any(s[2] for s in states))
        if not flagged:
            continue
        # A demo label anywhere on a duplicated identity vetoes supplied exposure.
        target = demo if any(_is_demo(r, dataset_kind) for r in records) else supplied
        amounts = defaultdict(set)
        unknown = False
        for record in records:
            invoice = record['invoice']
            amount = contracts.number(invoice.get('amount'))
            currency = invoice.get('currency')
            if amount is None or not isinstance(currency, str) or not re.fullmatch('[A-Z]{3}', currency):
                unknown = True
                continue
            currencies.add(currency)
            amounts[currency].add(amount)
        metrics['unknown_amount_flagged'] += int(unknown)
        for currency, values in amounts.items():
            if len(values) != 1:
                raise ValueError('CONFLICTING_BUSINESS_INVOICE_AMOUNT')
            target[currency].append(next(iter(values)))
    numbers = [value for bucket in (supplied, demo) for values in bucket.values() for value in values]
    with localcontext() as context:
        context.prec = _precision(numbers, len(businesses))
        metrics['supplied_flagged_exposure_by_currency'] = {
            currency: format(sum(supplied[currency], Decimal(0)), 'f') for currency in sorted(currencies)}
        metrics['demo_flagged_exposure_by_currency'] = {
            currency: format(sum(values, Decimal(0)), 'f') for currency, values in sorted(demo.items())}
    metrics['flagged_exposure_by_currency'] = dict(metrics['supplied_flagged_exposure_by_currency'])
    return metrics


def _duration(value):
    number = contracts.number(value)
    return number if number is not None and number >= 0 else None


def _timing(report):
    processing, review = (_duration(report.get(key)) for key in ('processing_seconds', 'review_seconds'))
    report.update(processing_seconds=report.get('processing_seconds') if processing is not None else None,
                  review_seconds=report.get('review_seconds') if review is not None else None,
                  processing_time_status='MEASURED' if processing is not None else 'NOT_MEASURED',
                  review_time_status='MEASURED' if review is not None else 'NOT_MEASURED')
    # Only the runner's explicit, user-reviewed policy provides cost assumptions.
    policy = report.get('policy', {})
    baseline = _duration(policy.get('baseline_minutes_per_invoice'))
    hourly = _duration(policy.get('hourly_cost'))
    currency = policy.get('cost_currency')
    total = report['metrics']['total']
    report.pop('cost_estimate', None)
    report['cost_status'] = 'NOT_CALCULATED'
    if (processing is None or review is None or baseline is None or hourly is None or not total
            or not isinstance(currency, str) or not re.fullmatch('[A-Z]{3}', currency)):
        return
    with localcontext() as context:
        context.prec = _precision([processing, review, baseline, hourly], total) * 2
        baseline_seconds = baseline * 60 * total
        measured_seconds = processing + review
        baseline_cost = baseline_seconds / 3600 * hourly
        measured_cost = measured_seconds / 3600 * hourly
        report['cost_estimate'] = {
            'currency': currency, 'baseline_minutes_per_invoice': baseline, 'hourly_cost': hourly,
            'baseline_seconds': baseline_seconds, 'measured_total_seconds': measured_seconds,
            'baseline_cost': baseline_cost.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
            'measured_time_cost': measured_cost.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
            'measured_time_cost_per_invoice': (measured_cost / total).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP),
            'estimated_cost_savings': (baseline_cost - measured_cost).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
            'method': 'User-supplied baseline minutes * 60 * unique invoices; measured total seconds = '
                      'processing + review; time cost = seconds / 3600 * hourly cost; per-invoice cost = '
                      'measured time cost / unique invoices (only when count > 0). Estimates, not prevented loss.'}
    report['cost_status'] = 'ESTIMATE_FROM_USER_BASELINE_AND_MEASURED_TIME'


def _display(value):
    if value is None:
        return 'UNKNOWN'
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)
    if isinstance(value, bool):
        return 'YES' if value else 'NO'
    return _text(value)


def _tables(report):
    tables = report.get('reference_tables', {})
    return tables if isinstance(tables, dict) else {}


def _reference_index(report):
    index = defaultdict(list)
    for table, records in _tables(report).items():
        if isinstance(records, list):
            for record in records:
                if isinstance(record, dict) and record.get('source_ref'):
                    index[str(record['source_ref'])].append((table, record))
    for result in report['results']:
        invoice = result['invoice']
        ref = invoice.get('source_ref')
        if ref and ('invoices', invoice) not in index[str(ref)]:
            index[str(ref)].append(('invoices', invoice))
    return index


def _evidence_rows(report, index):
    rows = []
    for result in report['results']:
        for domain, finding in _findings(result):
            for ref in finding.get('source_refs', []) or [None]:
                for table, record in index.get(str(ref), []) or [(None, None)]:
                    rows.append([result['business_id'], result['invoice'].get('invoice_id'), domain,
                                 finding.get('rule_id'), ref, table,
                                 'AVAILABLE' if record is not None else 'SOURCE_ROW_NOT_AVAILABLE', record])
    for source in report.get('source_manifest', []):
        rows.append([None, None, None, 'SOURCE_MANIFEST', source.get('source_ref') if isinstance(source, dict) else None,
                     'source_manifest', 'MANIFEST_ONLY', source])
    return rows


def _demo_rows(report):
    expected = report.get('demo_expected') or report.get('policy', {}).get('demo_expected', {})
    rows = []
    for invoice_id, rules in expected.items():
        matches = [r for r in report['results'] if _is_demo(r, report['dataset_kind'])
                   and str(r['invoice'].get('invoice_id')) == str(invoice_id)]
        for rule, status in rules.items():
            actual = sorted({f.get('status', 'UNKNOWN') for r in matches for _, f in _findings(r)
                             if f.get('rule_id') == rule})
            rows.append([invoice_id, rule, status, ', '.join(actual) if actual else 'UNKNOWN',
                         'PASS' if actual == [status] else 'UNKNOWN' if not actual else 'FAIL',
                         'SYNTHETIC_DEMO; excluded from supplied headline exposure'])
    return rows


def _vendor(result, report):
    invoice = result['invoice']
    vendors = _tables(report).get('vendors', [])
    matches = [row for row in vendors if isinstance(row, dict)
               and row.get('vendor_id') == invoice.get('vendor_id')] if isinstance(vendors, list) else []
    return (matches[0].get('vendor_name') or matches[0].get('name')) if len(matches) == 1 and (
        matches[0].get('vendor_name') or matches[0].get('name')) else invoice.get('vendor_name') or invoice.get('vendor_id')


def _summary_rows(report):
    metrics = report['metrics']
    rows = [[key.replace('_', ' ').title(), metrics[key], 'Unique business invoices', 'FROZEN']
            for key in ('total', 'covered', 'flagged', 'unresolved', 'unknown_amount_flagged')]
    for key, label in (('supplied_flagged_exposure_by_currency', 'Supplied flagged exposure (headline)'),
                       ('demo_flagged_exposure_by_currency', 'Synthetic demo flagged exposure (separate)')):
        for currency, amount in metrics[key].items():
            rows.append([label, contracts.number(amount), currency, 'FROZEN'])
    if not metrics['supplied_flagged_exposure_by_currency']:
        rows.append(['Supplied flagged exposure (headline)', Decimal(0), 'No supplied flagged currency', 'FROZEN'])
    for key in ('processing_seconds', 'review_seconds'):
        rows.append([key.replace('_', ' ').title(), contracts.number(report[key]), 'seconds',
                     report['processing_time_status' if key.startswith('processing') else 'review_time_status']])
    rows += [[key.replace('_', ' ').title(), report.get(key), '', 'SNAPSHOT']
             for key in ('report_id', 'version', 'created_at', 'dataset_kind', 'dataset_id', 'batch_id')]
    rows.append(['Cost status', report['cost_status'], '', 'SNAPSHOT'])
    for key, value in report.get('cost_estimate', {}).items():
        number = contracts.number(value)
        rows.append([key.replace('_', ' ').title(), number if number is not None else value,
                     report['cost_estimate']['currency'], 'ESTIMATE'])
    rows += [['Interpretation', DISCLAIMER, '', 'IMPORTANT'],
             ['Snapshot method', 'Executable Decimal calculations; values frozen; no Excel formulas. '
              'Covered means all four domain envelopes contain findings and no processing errors, not all checks PASS.', '', 'FROZEN'],
             ['Exposure method', 'One flagged business invoice once per currency; no domain/line sums or FX netting. '
              'Unknown amount or currency is excluded from exposure, not assumed zero. Demo labels veto supplied exposure.', '', 'FROZEN'],
             ['Processing errors', report.get('processing_errors', []), '', 'REVIEW_REQUIRED']]
    return rows


def _workbook(report, path, evidence, demo):
    # Import only; no installation or remote dependency resolution.
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def xl_color(css_hex):
        return 'FF' + css_hex.removeprefix('#').upper()

    XL_PRIMARY = xl_color('#2F5597')
    XL_HEADER = xl_color('#D9E2F3')
    XL_LIGHT = xl_color('#FAFCFF')
    XL_BORDER = xl_color('#D6DFEB')
    XL_RED = xl_color('#FFC7CE')
    XL_AMBER = xl_color('#FFEB9C')
    XL_GREEN = xl_color('#C6EFCE')
    status_colors = {'FAIL': XL_RED, 'UNKNOWN': XL_AMBER, 'NOT_MEASURED': XL_AMBER,
                     'SOURCE_ROW_NOT_AVAILABLE': XL_AMBER, 'PASS': XL_GREEN}
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = 'Invoice assurance approval snapshot'
    workbook.properties.description = DISCLAIMER
    invoice_rows, finding_rows = [], []
    for result in report['results']:
        invoice = result['invoice']
        amount = contracts.number(invoice.get('amount'))
        invoice_rows.append([result['business_id'], invoice.get('invoice_id'), invoice.get('invoice_number'),
                             invoice.get('invoice_date'), invoice.get('vendor_id'), _vendor(result, report),
                             'SYNTHETIC_DEMO' if _is_demo(result, report['dataset_kind']) else 'SUPPLIED',
                             result['status'], result['flagged'], result['unresolved'], amount,
                             invoice.get('currency'), format(amount, 'f') if amount is not None else None,
                             'COMPLETE' if _classification(result)[0] else 'UNKNOWN', invoice.get('source_ref')])
        for domain, finding in _findings(result):
            finding_rows.append([result['business_id'], invoice.get('invoice_id'), domain,
                                 finding.get('rule_id'), finding.get('status'), finding.get('severity'),
                                 finding.get('evidence_class'), finding.get('source_refs'), finding.get('observed_value'),
                                 finding.get('expected_value'), finding.get('explanation'), finding.get('recommended_action')])
    specs = [
        (SHEETS[0], ['Metric / metadata', 'Frozen value', 'Unit / currency', 'Measurement / scope'], _summary_rows(report), [44, 60, 28, 36]),
        (SHEETS[1], ['Business ID', 'Invoice ID', 'Invoice number', 'Invoice date', 'Vendor ID', 'Vendor', 'Dataset kind',
                    'Status', 'Flagged', 'Unresolved', 'Amount', 'Currency', 'Amount exact', 'Coverage', 'Source reference'],
         invoice_rows, [38, 20, 22, 18, 20, 30, 24, 20, 14, 16, 24, 14, 36, 20, 40]),
        (SHEETS[2], ['Business ID', 'Invoice ID', 'Domain', 'Rule ID', 'Status', 'Severity', 'Evidence class',
                    'Source references', 'Observed value', 'Expected value', 'Explanation', 'Recommended action'],
         finding_rows, [38, 20, 20, 30, 20, 16, 26, 44, 50, 50, 58, 58]),
        (SHEETS[3], ['Business ID', 'Invoice ID', 'Domain', 'Rule ID', 'Source reference', 'Source table', 'Availability', 'Sanitized source record'],
         evidence, [38, 20, 20, 30, 44, 28, 34, 90]),
        (SHEETS[4], ['Invoice ID', 'Rule ID', 'Expected status', 'Actual status', 'Comparison status', 'Scope'],
         demo, [22, 34, 24, 28, 26, 68]),
    ]
    for name, headers, rows, widths in specs:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = XL_PRIMARY
        sheet.freeze_panes = 'B4'
        sheet.cell(1, 1, name + ' | approval snapshot')
        sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color=XL_PRIMARY)
        sheet.cell(2, 1, 'Static values • no formulas • explicit approval required')
        sheet['A2'].font = Font(name='Calibri', size=11, color=XL_PRIMARY)
        header_row, data_start, data_end = 3, 4, 3 + len(rows)
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(header_row, column, header)
            cell.font = Font(name='Calibri', bold=True, color=XL_PRIMARY)
            cell.fill = PatternFill('solid', fgColor=XL_HEADER)
            cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
            cell.border = Border(bottom=Side(style='thin', color=XL_BORDER))
            cell.comment = Comment('Frozen snapshot. Statuses: PASS, FAIL, UNKNOWN, NOT_APPLICABLE. '
                                   'Missing measurements: NOT_MEASURED.', 'Invoice assurance')
        for row_number, row in enumerate(rows, data_start):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row_number, column)
                exact_decimal = isinstance(value, Decimal)
                if value is None:
                    value = 'NOT_MEASURED' if name == 'Summary' and row[3] == 'NOT_MEASURED' else 'UNKNOWN'
                elif isinstance(value, (dict, list)):
                    value = _display(value)
                elif exact_decimal:
                    # Excel has only 15 significant digits; preserve larger exact values as text.
                    if len(value.as_tuple().digits) > 15 or not -300 <= value.adjusted() <= 300:
                        value = format(value, 'f')
                        cell.comment = Comment('Exact Decimal text: exceeds Excel numeric precision/range.', 'Invoice assurance')
                if isinstance(value, str):
                    cell.value = _text(value)[:32767]
                    cell.data_type = 's'  # All source strings, including = + - @, are literal text.
                    cell.number_format = '@'
                else:
                    cell.value = value
                    cell.number_format = '#,##0.00########' if exact_decimal else '0'
                cell.font = Font(name='Calibri', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True,
                                           horizontal='right' if exact_decimal else 'left')
                cell.fill = PatternFill('solid', fgColor=status_colors.get(str(value), XL_LIGHT))
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{max(header_row, data_end)}'
        sheet.print_title_rows = '1:3'
        sheet.page_setup.orientation = 'landscape'
    # Exclusive creation plus a fresh directory prevents overwrite on every path.
    with path.open('xb') as stream:
        workbook.save(stream)
    workbook.close()


def _html_table(headers, rows):
    return '<div class="table-wrap"><table><thead><tr>' + ''.join(
        '<th>' + escape(_display(header)) + '</th>' for header in headers) + '</tr></thead><tbody>' + ''.join(
        '<tr>' + ''.join('<td>' + escape(_display(value)) + '</td>' for value in row) + '</tr>' for row in rows
    ) + '</tbody></table></div>'


def _html(report, index, demo):
    metrics = report['metrics']
    cards = ''.join('<div class="card"><span>' + escape(label) + '</span><strong>' + str(metrics[key]) + '</strong></div>'
                    for key, label in (('total', 'Unique invoices'), ('covered', 'Domain coverage'),
                                       ('flagged', 'Flagged invoices'), ('unresolved', 'Unresolved invoices')))
    exposure_rows = [[currency, metrics['supplied_flagged_exposure_by_currency'].get(currency, '0'),
                      metrics['demo_flagged_exposure_by_currency'].get(currency, '0')]
                     for currency in sorted(set(metrics['supplied_flagged_exposure_by_currency']) |
                                            set(metrics['demo_flagged_exposure_by_currency']))]
    exposure = _html_table(['Currency', 'Supplied headline exposure', 'Synthetic demo exposure (separate)'],
                          exposure_rows or [['No known flagged currency', '0', '0']])
    invoices = []
    for result in report['results']:
        invoice = result['invoice']
        title = ' · '.join(_display(value) for value in (invoice.get('invoice_id') or invoice.get('invoice_number'),
                           _vendor(result, report), invoice.get('amount'), invoice.get('currency'), result['status']))
        body = '<p class="scope">' + escape(result['business_id']) + ' · ' + (
            'SYNTHETIC_DEMO — excluded from supplied headline exposure' if _is_demo(result, report['dataset_kind']) else 'SUPPLIED') + '</p>'
        body += _html_table(['Invoice field', 'Sanitized value'], invoice.items())
        lines = _tables(report).get('invoice_lines')
        if isinstance(lines, list):
            linked = [row for row in lines if isinstance(row, dict) and invoice.get('invoice_id') is not None
                      and str(row.get('invoice_id')) == str(invoice['invoice_id'])]
            body += '<h3>Exact invoice line records</h3>'
            body += _html_table(['Source reference', 'Sanitized record'], [[row.get('source_ref'), row] for row in linked]) if linked else '<p>UNKNOWN — sourced invoice lines are missing; header-only matching is not sufficient.</p>'
        for domain in contracts.DOMAINS:
            envelope = result.get('domains', {}).get(domain)
            domain_body = ''
            if not isinstance(envelope, dict) or not envelope.get('findings'):
                domain_body = '<p>UNKNOWN — domain findings are missing.</p>'
            else:
                for finding in envelope['findings']:
                    fields = [(label, finding.get(key)) for label, key in (
                        ('Rule', 'rule_id'), ('Status', 'status'), ('Evidence class', 'evidence_class'),
                        ('Severity', 'severity'), ('Observed', 'observed_value'), ('Expected', 'expected_value'),
                        ('Explanation', 'explanation'), ('Recommended action', 'recommended_action'))]
                    sources = [[ref, table, record] for ref in finding.get('source_refs', []) or [None]
                               for table, record in index.get(str(ref), []) or [('SOURCE_ROW_NOT_AVAILABLE', None)]]
                    domain_body += '<details class="finding"><summary>' + escape(_display(finding.get('rule_id'))) + ' · ' + escape(_display(finding.get('status'))) + '</summary>'
                    domain_body += _html_table(['Finding field', 'Value'], fields)
                    domain_body += '<h4>Exact source references and records</h4>' + _html_table(
                        ['Source reference', 'Source table / availability', 'Sanitized record'], sources)
                    domain_body += '</details>'
                if envelope.get('errors'):
                    domain_body += '<p>Domain processing errors: ' + escape(_display(envelope['errors'])) + '</p>'
            body += '<details class="domain"><summary>' + escape(domain.title()) + '</summary>' + domain_body + '</details>'
        invoices.append('<details class="invoice"><summary>' + escape(title) + '</summary>' + body + '</details>')
    title = 'Invoice assurance · ' + _display(report.get('report_id')) + ' · v' + _display(report.get('version'))
    return '''<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'">
<title>''' + escape(title) + '''</title><style>
:root{color-scheme:light;font-family:system-ui,-apple-system,Segoe UI,sans-serif;color:#24344b;background:#f5f8fc}
*{box-sizing:border-box}body{margin:0}main{max-width:1280px;margin:auto;padding:40px 28px 64px}
h1{font-size:32px;letter-spacing:-.8px;margin:10px 0}h2{font-size:22px;margin:30px 0 14px}h3,h4{font-size:16px}
p{line-height:1.6}.eyebrow,.scope{color:#536986;font-size:13px;overflow-wrap:anywhere}.eyebrow{letter-spacing:1px;text-transform:uppercase}
.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:16px;margin:24px 0}.card{background:white;border:1px solid #d6dfeb;border-radius:12px;padding:22px}
.card span{display:block;color:#536986;font-size:14px}.card strong{display:block;font-size:36px;margin-top:8px;color:#2f5597}
.notice{background:#fff5db;border-left:4px solid #b67b0e;border-radius:6px;padding:14px 18px}
details{border:1px solid #d6dfeb;border-radius:10px;background:white;margin:12px 0;padding:0 18px 16px}
summary{cursor:pointer;padding:16px 0;font-weight:650;overflow-wrap:anywhere}details:not([open]){padding-bottom:0}
.domain{background:#fafcff}.finding{border-radius:6px}.table-wrap{overflow-x:auto;margin:12px 0}
table{width:100%;border-collapse:collapse;font-size:13px}th,td{text-align:left;vertical-align:top;border-bottom:1px solid #d6dfeb;padding:12px;overflow-wrap:anywhere;white-space:pre-wrap}
th{background:#eaf0f9;color:#2f5597}td{min-width:140px;max-width:560px}footer{color:#536986;font-size:13px;margin-top:28px}
@media(max-width:700px){main{padding:24px 16px}.cards{grid-template-columns:repeat(2,1fr)}h1{font-size:26px}}
</style></head><body><main><div class="eyebrow">Frozen approval snapshot · manual review required</div><h1>''' + escape(title) + '</h1><p class="scope">' + escape(_display(report.get('created_at'))) + ' · ' + escape(report['dataset_kind']) + '</p><p class="notice">' + escape(DISCLAIMER) + '</p><section class="cards">' + cards + '</section><h2>Flagged exposure by currency</h2>' + exposure + '<p>Unknown amount / currency flagged invoices: ' + str(metrics['unknown_amount_flagged']) + '. No currency conversion or domain double counting.</p><details><summary>Snapshot metrics, measured timing and methodology</summary>' + _html_table(['Metric', 'Frozen value', 'Unit', 'Status'], _summary_rows(report)) + '</details><h2>Invoice drilldown</h2>' + ''.join(invoices) + '<details><summary>Source manifest</summary>' + _html_table(['Source', 'Sanitized manifest record'], [[i, row] for i, row in enumerate(report.get('source_manifest', []), 1)]) + '</details><details><summary>Processing errors — preserved for review</summary><p>' + escape(_display(report.get('processing_errors', []))) + '</p></details><details><summary>Synthetic demo expected-versus-actual cases</summary>' + _html_table(['Invoice', 'Rule', 'Expected', 'Actual', 'Comparison', 'Scope'], demo) + '</details><footer>Executable Decimal calculations frozen at generation. No live formulas, external assets or automatic clearance. Missing measurements are NOT_MEASURED.</footer></main></body></html>'


def generate(report, output_dir):
    """Write dashboard, workbook and sanitized JSON into a never-used directory."""
    folder = Path(output_dir).expanduser().resolve()
    if folder.exists():
        raise FileExistsError('SNAPSHOT_DIRECTORY_ALREADY_EXISTS')
    # Compute from original evidence before redaction, never from supplied summary numbers.
    snapshot = deepcopy(report)
    snapshot['metrics'] = summarize(report['results'], report['dataset_kind'])
    _timing(snapshot)
    snapshot = _sanitize(snapshot)
    for result in snapshot['results']:
        _, flagged, unresolved = _classification(result)
        result.update(flagged=flagged, unresolved=unresolved,
                      status='FAIL' if flagged else 'UNKNOWN' if unresolved else 'PASS')
    index = _reference_index(snapshot)
    evidence, demo = _evidence_rows(snapshot, index), _demo_rows(snapshot)
    html = _html(snapshot, index, demo)
    data = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)
    paths = {key: folder / filename for key, filename in (
        ('dashboard', 'dashboard.html'), ('workbook', 'assurance.xlsx'), ('data', 'report.json'))}
    folder.mkdir(parents=True, exist_ok=False)
    with paths['dashboard'].open('x', encoding='utf-8') as stream:
        stream.write(html)
    with paths['data'].open('x', encoding='utf-8') as stream:
        stream.write(data)
    _workbook(snapshot, paths['workbook'], evidence, demo)
    return {key: str(path) for key, path in paths.items()}
